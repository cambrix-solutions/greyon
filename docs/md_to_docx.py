from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

md_path = Path(__file__).resolve().parent / "BACKEND_API.md"
out_path = Path(__file__).resolve().parent / "Greyon_Backend_API.docx"
text = md_path.read_text(encoding="utf-8")

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def set_code_shading(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F4F6")
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)


def add_runs_with_inline(paragraph, content):
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")
    parts = pattern.split(content)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x33, 0x3A, 0x40)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def add_table(rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            cell.text = row[j] if j < len(row) else ""
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True
    doc.add_paragraph()


lines = text.splitlines()
i = 0
in_code = False
code_buf = []
table_buf = []


def flush_table():
    global table_buf
    if not table_buf:
        return
    rows = []
    for raw in table_buf:
        cells = [c.strip() for c in raw.strip().strip("|").split("|")]
        if all(
            re.fullmatch(r":?-+:?", c.replace(" ", "")) for c in cells if c
        ):
            continue
        rows.append(cells)
    add_table(rows)
    table_buf = []


while i < len(lines):
    line = lines[i]

    if line.strip().startswith("```"):
        if not in_code:
            flush_table()
            in_code = True
            code_buf = []
        else:
            in_code = False
            p = doc.add_paragraph()
            set_code_shading(p)
            run = p.add_run("\n".join(code_buf))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            code_buf = []
        i += 1
        continue

    if in_code:
        code_buf.append(line)
        i += 1
        continue

    if line.strip().startswith("|"):
        table_buf.append(line)
        i += 1
        continue
    else:
        flush_table()

    if not line.strip():
        i += 1
        continue

    m = re.match(r"^(#{1,6})\s+(.*)$", line)
    if m:
        level = len(m.group(1))
        title = re.sub(r"\{#.*\}$", "", m.group(2)).strip()
        if level == 1:
            doc.add_heading(title, level=0)
        else:
            doc.add_heading(title, level=min(level, 3))
        i += 1
        continue

    if line.startswith(">"):
        content = line.lstrip("> ").strip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        run_prefix = p.add_run("Note: ")
        run_prefix.bold = True
        run_prefix.italic = True
        add_runs_with_inline(p, content)
        for r in p.runs:
            r.italic = True
        i += 1
        continue

    if re.match(r"^[-*]\s+", line) or re.match(r"^\d+\.\s+", line):
        content = re.sub(r"^([-*]|\d+\.)\s+", "", line)
        style_name = "List Bullet" if re.match(r"^[-*]\s+", line) else "List Number"
        p = doc.add_paragraph(style=style_name)
        add_runs_with_inline(p, content)
        i += 1
        continue

    if re.match(r"^-{3,}$", line.strip()) or re.match(r"^\*{3,}$", line.strip()):
        i += 1
        continue

    p = doc.add_paragraph()
    add_runs_with_inline(p, line.strip())
    i += 1

flush_table()
doc.save(out_path)
print(f"Wrote {out_path} ({out_path.stat().st_size} bytes)")
